"""Collections (CRUD, study stream, game, chat) and the media/notes listing endpoints."""
import io
import json
import re
import threading
import time
import uuid
from pathlib import Path

import pytest

import agents.media_agent
from tests.conftest import make_pdf


def _events(raw: str):
    return [(f.split("\n")[0].split(": ", 1)[1], json.loads(f.split("data: ", 1)[1]))
            for f in raw.strip().split("\n\n")]


@pytest.fixture(autouse=True)
def own_user(client):
    """The data dir is shared across tests; listings here need a user of their own."""
    client.headers["X-User-Id"] = f"u-{uuid.uuid4().hex[:12]}"


@pytest.fixture
def lecture(client, monkeypatch):
    monkeypatch.setattr(agents.media_agent, "transcribe_file", lambda path, **_: "Photosynthesis makes sugar.")
    r = client.post("/media/lectures", data={"title": "Bio 101"},
                    files={"audio": ("lecture.wav", b"RIFF0000WAVEfmt ", "audio/wav")})
    assert r.status_code == 200, r.text
    return r.json()


@pytest.fixture
def presentation(client):
    from pptx import Presentation
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Cells"
    slide.placeholders[1].text = "Cells are the unit of life"
    buf = io.BytesIO()
    prs.save(buf)
    r = client.post("/media/presentations",
                    files={"presentation": ("deck.pptx", buf.getvalue(),
                                            "application/vnd.openxmlformats-officedocument.presentationml.presentation")})
    assert r.status_code == 200, r.text
    return r.json()


def _poll(client, task_id: str, until=lambda task: task["status"] in ("done", "failed")) -> dict:
    for _ in range(500):
        task = client.get(f"/media/tasks/{task_id}").json()
        if until(task):
            return task
        time.sleep(0.01)
    raise AssertionError(f"task never got there: {task}")


def test_youtube_transcribes_in_the_background_with_progress(client, monkeypatch):
    gate = threading.Event()  # holds the transcription so the test sees it in progress

    def download(url, dest, on_progress):
        for n in range(1, 5):
            on_progress(n / 4, "Cell Biology")
        path = Path(dest) / "abc.m4a"
        path.write_bytes(b"audio")
        return path, "Cell Biology", "abc"

    def transcribe(path, on_progress):
        on_progress(0.5)
        gate.wait(5)
        on_progress(1.0)
        return "Cells divide by mitosis."

    monkeypatch.setattr(agents.media_agent, "download_youtube_audio", download)
    monkeypatch.setattr(agents.media_agent, "transcribe_file", transcribe)
    r = client.post("/media/youtube", json={"video_url": "https://youtu.be/abc"})
    assert r.status_code == 202 and r.json()["status"] in ("queued", "running")
    task_id = r.json()["task_id"]

    task = _poll(client, task_id, until=lambda t: t["stage"] == "transcribing" and t["progress"] == 0.5)
    assert task["status"] == "running" and task["title"] == "Cell Biology" and task["result"] is None
    assert client.get(f"/media/tasks/{task_id}", headers={"X-User-Id": "someone-else"}).status_code == 404

    gate.set()
    done = _poll(client, task_id)
    assert done["status"] == "done" and done["progress"] == 1.0 and not done["error"]
    assert done["result"]["transcript"] == "Cells divide by mitosis." and done["result"]["title"] == "Cell Biology"
    assert [t["job_id"] for t in client.get("/media/transcriptions").json()] == [done["result"]["job_id"]]


def test_youtube_task_reports_failure(client, monkeypatch):
    def download(url, dest, on_progress):
        raise RuntimeError("Video unavailable")

    monkeypatch.setattr(agents.media_agent, "download_youtube_audio", download)
    task_id = client.post("/media/youtube", json={"video_url": "https://youtu.be/gone"}).json()["task_id"]
    failed = _poll(client, task_id)
    assert failed["status"] == "failed" and failed["error"] == "Video unavailable" and failed["result"] is None
    assert client.get("/media/tasks/nope").status_code == 404


def test_media_lists_and_details(client, lecture, presentation):
    listed = client.get("/media/transcriptions").json()
    assert [t["job_id"] for t in listed] == [lecture["job_id"]]
    assert listed[0]["collection_id"] == lecture["collection_id"]

    detail = client.get(f"/media/transcriptions/{lecture['job_id']}").json()
    assert detail["transcript"] == "Photosynthesis makes sugar." and detail["metadata"]["title"] == "Bio 101"
    assert client.get("/media/transcriptions/nope").status_code == 404

    pres = client.get("/media/presentations").json()
    assert pres[0]["presentation_id"] == presentation["presentation_id"]
    assert pres[0]["collection_id"] == presentation["collection_id"]
    slides = client.get(f"/media/presentations/{presentation['presentation_id']}").json()["slides"]
    assert slides[0]["title"] == "Cells"


def test_notes_listing(client):
    r = client.post("/notes", files={"notes": ("notes.pdf", make_pdf(4, outline=False), "application/pdf")})
    assert r.status_code == 200, r.text
    listed = client.get("/notes").json()
    assert [n["notes_id"] for n in listed] == [r.json()["notes_id"]]
    assert listed[0]["collection_id"] == r.json()["collection_id"]


def test_auto_collections_are_hidden_but_studyable(client, lecture, llm):
    assert client.get("/collections").json() == []
    auto = client.get("/collections", params={"include_auto": True}).json()
    assert [c["collection_id"] for c in auto] == [lecture["collection_id"]]

    r = client.post(f"/collections/{lecture['collection_id']}/study", json={})
    assert r.status_code == 200 and r.json()["narrative"]


def test_collection_lifecycle(client, book, lecture, llm):
    col = client.post("/collections", json={"name": "Mix", "materials": {
        "textbook_sections": [{"file_id": book["file_id"], "section_id": "ch1.s1", "title": "One"}],
    }}).json()
    cid = col["collection_id"]
    assert col["materials"]["transcriptions"] == []  # every kind is present
    assert [c["collection_id"] for c in client.get("/collections").json()] == [cid]

    assert client.get(f"/collections/{cid}/study").status_code == 404
    with client.stream("POST", f"/collections/{cid}/study/stream", json={}) as r:
        assert r.headers["cache-control"] == "no-cache, no-transform"
        events = _events(r.read().decode())
    assert events[-1][0] == "done" and any(e == "token" for e, _ in events)
    assert client.get(f"/collections/{cid}/study").json() == events[-1][1]

    game = client.post(f"/collections/{cid}/game")
    assert game.status_code == 200 and game.json()["game_code"]
    fix = client.post(f"/collections/{cid}/game/fix",
                      json={"error": "TypeError: x is undefined", "version": game.json()["game_version"]})
    assert fix.status_code == 200 and fix.json()["game_version"] == game.json()["game_version"] + 1

    chat = client.post(f"/collections/{cid}/chat", json={"message": "hi"})
    assert chat.status_code == 200
    assert len(client.get(f"/collections/{cid}/chat").json()["history"]) == 2

    assert client.patch(f"/collections/{cid}", json={"name": "Renamed"}).json()["name"] == "Renamed"
    assert client.patch(f"/collections/{cid}", json={"name": " "}).status_code == 400

    # changing the contents drops the generated materials, not the chat
    materials = {**col["materials"], "transcriptions": [{"transcription_id": lecture["job_id"]}]}
    r = client.put(f"/collections/{cid}/materials", json=materials)
    assert r.status_code == 200 and len(r.json()["materials"]["transcriptions"]) == 1
    assert client.get(f"/collections/{cid}/study").status_code == 404
    assert len(client.get(f"/collections/{cid}/chat").json()["history"]) == 2

    assert client.delete(f"/collections/{cid}").status_code == 204
    assert client.get(f"/collections/{cid}").status_code == 404
    assert client.get(f"/collections/{cid}/chat").status_code == 404
    assert client.delete(f"/collections/{cid}").status_code == 404


def test_notes_auto_collection_feeds_its_notes(client, llm):
    notes = client.post("/notes", files={"notes": ("notes.pdf", make_pdf(4, outline=False), "application/pdf")}).json()
    cid = notes["collection_id"]
    refs = client.get(f"/collections/{cid}").json()["materials"]["notes"]
    assert [r["notes_id"] for r in refs] == [notes["notes_id"]]

    assert client.post(f"/collections/{cid}/study", json={}).status_code == 200
    assert any("Page 3 content about topic 3" in prompt for prompt in llm.calls)


def test_collections_saved_with_note_id_still_load(client):
    from services.storage_service import save_collection
    user = client.headers["X-User-Id"]
    save_collection(user, "legacy", {"collection_id": "legacy", "name": "Old", "user_id": user, "auto": True,
                                     "materials": {"notes": [{"note_id": "n1"}]}})
    assert client.get("/collections/legacy").json()["materials"]["notes"] == [{"notes_id": "n1"}]


def test_collection_tutor_searches_only_its_sources(client, book, lecture, llm):
    col = client.post("/collections", json={"name": "Lecture only", "materials": {
        "transcriptions": [{"transcription_id": lecture["job_id"]}]}}).json()
    llm.tool_calls = [[("search_materials", {"query": "photosynthesis", "max_results": 8})]]
    assert client.post(f"/collections/{col['collection_id']}/chat", json={"message": "hi"}).status_code == 200

    found = llm.conversations[-1][-1].content
    assert re.findall(r'source="([^"]+)"', found) == [f"transcriptions/{lecture['job_id']}"]
