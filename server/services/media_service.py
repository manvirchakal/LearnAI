"""
Media processing service — YouTube + lecture audio transcription via faster-whisper,
PowerPoint processing via python-pptx.
Replaces AWS Transcribe and AWS S3 audio uploads.
"""
import asyncio
import logging
import os
import subprocess
import tempfile
import uuid
from datetime import datetime
from typing import Optional

from yt_dlp import YoutubeDL
from faster_whisper import WhisperModel
from pptx import Presentation
from functools import lru_cache

from core.config import settings
from services.storage_service import (
    save_transcription_metadata,
    save_transcription_content,
    save_presentation_metadata,
    save_presentation_slide,
)

logger = logging.getLogger(__name__)

SUPPORTED_AUDIO_FORMATS = [".mp3", ".wav", ".m4a"]
SUPPORTED_VIDEO_FORMATS = [".mp4"]
SUPPORTED_MEDIA_FORMATS = SUPPORTED_AUDIO_FORMATS + SUPPORTED_VIDEO_FORMATS


@lru_cache(maxsize=1)
def _whisper_model() -> WhisperModel:
    logger.info(f"Loading Whisper model: {settings.WHISPER_MODEL}")
    return WhisperModel(settings.WHISPER_MODEL, device="cpu", compute_type="int8")


def _transcribe_file(audio_path: str) -> str:
    """Run faster-whisper transcription on a local audio file."""
    model = _whisper_model()
    segments, _ = model.transcribe(audio_path, beam_size=5)
    return " ".join(seg.text.strip() for seg in segments)


def _convert_to_mp3(input_path: str, output_path: str) -> None:
    """Convert audio/video to MP3 using ffmpeg."""
    result = subprocess.run(
        ["ffmpeg", "-i", input_path, "-vn", "-acodec", "libmp3lame", "-ab", "192k",
         "-ar", "44100", "-y", output_path],
        check=True,
        capture_output=True,
    )


async def transcribe_youtube(video_url: str, user_id: str) -> dict:
    """Download YouTube audio and transcribe with local Whisper."""
    job_id = str(uuid.uuid4())
    temp_audio_path = f"/tmp/{job_id}"

    ydl_opts = {
        "format": "bestaudio/best",
        "postprocessors": [{"key": "FFmpegExtractAudio", "preferredcodec": "mp3", "preferredquality": "192"}],
        "outtmpl": temp_audio_path,
        "quiet": True,
        "no_warnings": True,
    }

    try:
        with YoutubeDL(ydl_opts) as ydl:
            info = ydl.extract_info(video_url, download=True)
            video_title = info.get("title", "Untitled Video")
            video_id = info.get("id", "")

        mp3_path = f"{temp_audio_path}.mp3"
        logger.info(f"Transcribing YouTube audio: {mp3_path}")
        transcript_text = await asyncio.get_event_loop().run_in_executor(
            None, _transcribe_file, mp3_path
        )

        metadata = {
            "job_id": job_id,
            "title": video_title,
            "original_filename": video_title,
            "video_id": video_id,
            "video_url": video_url,
            "transcription_date": datetime.now().isoformat(),
            "source_type": "youtube",
        }
        save_transcription_metadata(user_id, job_id, metadata)
        save_transcription_content(user_id, job_id, transcript_text)

        from services.collection_service import create_default_collection
        collection_id = create_default_collection("transcriptions", job_id, metadata, user_id)

        return {
            "transcript": transcript_text,
            "metadata": metadata,
            "job_id": job_id,
            "collection_id": collection_id,
            "video_title": video_title,
        }
    finally:
        for ext in [".mp3", ".m4a", ".webm", ""]:
            path = f"{temp_audio_path}{ext}"
            if os.path.exists(path):
                os.remove(path)


async def transcribe_lecture(audio_bytes: bytes, filename: str, title: str, user_id: str) -> dict:
    """Transcribe uploaded lecture audio/video using local Whisper."""
    job_id = str(uuid.uuid4())
    file_ext = os.path.splitext(filename)[1].lower()

    if file_ext not in SUPPORTED_MEDIA_FORMATS:
        raise ValueError(f"Unsupported format {file_ext}. Supported: {SUPPORTED_MEDIA_FORMATS}")

    temp_input = f"/tmp/{job_id}_input{file_ext}"
    temp_mp3 = f"/tmp/{job_id}.mp3"

    try:
        with open(temp_input, "wb") as f:
            f.write(audio_bytes)

        if file_ext != ".mp3":
            _convert_to_mp3(temp_input, temp_mp3)
            transcribe_path = temp_mp3
        else:
            transcribe_path = temp_input

        transcript_text = await asyncio.get_event_loop().run_in_executor(
            None, _transcribe_file, transcribe_path
        )

        metadata = {
            "job_id": job_id,
            "title": title,
            "original_filename": title,
            "file_name": filename,
            "upload_date": datetime.now().isoformat(),
            "file_type": "uploaded_audio" if file_ext in SUPPORTED_AUDIO_FORMATS else "uploaded_video",
            "source_type": "upload",
        }
        save_transcription_metadata(user_id, job_id, metadata)
        save_transcription_content(user_id, job_id, transcript_text)

        from services.collection_service import create_default_collection
        collection_id = create_default_collection("transcriptions", job_id, metadata, user_id)

        return {"transcript": transcript_text, "metadata": metadata, "job_id": job_id, "collection_id": collection_id}
    finally:
        for path in [temp_input, temp_mp3]:
            if os.path.exists(path):
                os.remove(path)


def process_presentation(pptx_bytes: bytes, filename: str, user_id: str) -> dict:
    """Parse a PowerPoint file and store slides locally."""
    presentation_id = str(uuid.uuid4())
    temp_path = f"/tmp/{presentation_id}.pptx"

    try:
        with open(temp_path, "wb") as f:
            f.write(pptx_bytes)

        prs = Presentation(temp_path)
        slides_content = []

        for num, slide in enumerate(prs.slides, 1):
            slide_data = {"number": num, "title": "", "content": [], "notes": ""}
            if slide.shapes.title:
                slide_data["title"] = slide.shapes.title.text
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_data["content"].append(shape.text.strip())
            if slide.notes_slide and slide.notes_slide.notes_text_frame:
                slide_data["notes"] = slide.notes_slide.notes_text_frame.text
            slides_content.append(slide_data)

        metadata = {
            "presentation_id": presentation_id,
            "original_filename": filename,
            "upload_date": datetime.now().isoformat(),
            "total_slides": len(slides_content),
            "has_speaker_notes": any(s["notes"] for s in slides_content),
            "slide_titles": [s["title"] for s in slides_content if s["title"]],
        }

        save_presentation_metadata(user_id, presentation_id, metadata)
        for slide in slides_content:
            save_presentation_slide(user_id, presentation_id, slide["number"],
                                    {"title": slide["title"], "content": slide["content"], "notes": slide["notes"]})

        from services.collection_service import create_default_collection
        collection_id = create_default_collection("presentations", presentation_id, metadata, user_id)

        return {"message": "Presentation processed successfully", "presentation_id": presentation_id,
                "metadata": metadata, "collection_id": collection_id}
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)
